#!/usr/bin/env python3
"""trees.json → yksi Word-dokumentti per sivu.

Muoto: juokseva lukudokumentti. Word-otsikkotasot, placeholder-luonnokset
keltaisella korostuksella, sisäiset KYSYMYS-notet dokumentin loppuliitteeseen.
"""
import io, os, re, json, sys, datetime, unicodedata

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, 'word-export')

GREY = RGBColor(0x6F, 0x6F, 0x6F)
DARK = RGBColor(0x1D, 0x1D, 0x1F)
MID = RGBColor(0x55, 0x55, 0x55)
BLUE = RGBColor(0x0A, 0x6E, 0x9B)

# Inline-kysymys copyn sisällä: "... teksti. [KYSYMYS 9: tarkka vuosi?]"
INLINE_Q = re.compile(r'\[\s*(?:KYSYMYS|Kysymys)\b[^\]]*\]')
# Hakasulkeissa oleva wireframe-luonnos
BRACKET = re.compile(r'\[[^\]]*\]')
NOTE_Q = re.compile(r'kysymys', re.I)


def slug(s):
    s = unicodedata.normalize('NFKD', s)
    s = ''.join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r'[^A-Za-z0-9]+', '_', s).strip('_')
    return s


# ---------------------------------------------------------------- tyylit
def setup_styles(doc):
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    st.font.color.rgb = DARK
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    pf = st.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15

    for name, size, color in (('Heading 1', 20, BLUE), ('Heading 2', 15, DARK),
                              ('Heading 3', 12, DARK), ('Heading 4', 11, MID)):
        s = doc.styles[name]
        s.font.name = 'Calibri'
        s.font.size = Pt(size)
        s.font.color.rgb = color
        s.font.bold = True
        s.paragraph_format.space_before = Pt(14 if size >= 15 else 10)
        s.paragraph_format.space_after = Pt(4)

    t = doc.styles['Title']
    t.font.name = 'Calibri'
    t.font.size = Pt(28)
    t.font.color.rgb = DARK
    t.font.bold = True


def para(doc, style=None):
    p = doc.add_paragraph(style=style)
    return p


def add_rich(p, text, base=None):
    """Lisää teksti niin, että [hakasulkeissa] oleva luonnos korostuu keltaisella."""
    pos = 0
    for m in BRACKET.finditer(text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
            if base:
                base(r)
        r = p.add_run(m.group(0))
        if base:
            base(r)
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        if base:
            base(r)
    if not text:
        p.add_run('')
    return p


def grey(sz=9, italic=True):
    def f(r):
        r.font.size = Pt(sz)
        r.font.color.rgb = GREY
        r.font.italic = italic
    return f


def lede_fmt(r):
    r.font.size = Pt(12)
    r.font.color.rgb = MID


def eyebrow_fmt(r):
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = BLUE


def bold_fmt(r):
    r.font.bold = True


def small_fmt(r):
    r.font.size = Pt(10)


STEP_NUM = re.compile(r'^(\d{1,2})\s+(\S.*)$')
STAT_VAL = re.compile(r'^[\d\s+%<>~,.\[\]A-ZÅÄÖ/-]{1,10}$')


def regroup_steps(items):
    """SVG-vaiheistuksessa laatikot ovat vuorotellen viivan ylä- ja alapuolella,
    joten DOM-järjestys on 1,3,5,2,4,6. Järjestä numeroidut vaiheet uudelleen."""
    out, i = [], 0
    while i < len(items):
        run, j = [], i
        while j < len(items):
            it = items[j]
            m = STEP_NUM.match(it['text']) if it['role'] in ('p', 'h3', 'h4', 'text') else None
            if not m:
                break
            desc = None
            nxt = items[j + 1] if j + 1 < len(items) else None
            if nxt and nxt['role'] in ('p', 'text', 'li') and not STEP_NUM.match(nxt['text']):
                desc = nxt
                j += 1
            run.append((int(m.group(1)), m.group(2), desc))
            j += 1
        if len(run) >= 3:
            run.sort(key=lambda x: x[0])
            for num, title, desc in run:
                out.append({'role': 'step', 'text': f'{num}. {title}', 'tag': 'step'})
                if desc:
                    out.append(desc)
            i = j
        else:
            out.append(items[i])
            i += 1
    return out


def pair_stats(items):
    """Lukunosto: arvo ja selite ovat eri elementeissä — yhdistä yhdeksi riviksi."""
    out, i = [], 0
    while i < len(items):
        a = items[i]
        b = items[i + 1] if i + 1 < len(items) else None
        if (a['role'] in ('text', 'p') and len(a['text']) <= 10 and STAT_VAL.match(a['text'])
                and b and b['role'] in ('text', 'p') and len(b['text']) <= 45
                and not STAT_VAL.match(b['text'])):
            out.append({'role': 'stat', 'text': a['text'] + ' — ' + b['text'], 'tag': 'stat'})
            i += 2
            continue
        out.append(a)
        i += 1
    return out


# ---------------------------------------------------------------- sisältö
def clean(text, questions, where):
    """Irrota inline-KYSYMYKSET liitteeseen ja palauta puhdas copy."""
    found = INLINE_Q.findall(text)
    for q in found:
        questions.append((where, q.strip('[]').strip()))
    if found:
        text = INLINE_Q.sub('', text)
    text = re.sub(r'\s{2,}', ' ', text)
    text = re.sub(r'\s+([.,:;!?])', r'\1', text)      # inline-liitos jättää " ." → "."
    return text.strip()


def render_items(doc, items, questions, where):
    items = pair_stats(regroup_steps(items))
    chips = []

    def flush_chips():
        if not chips:
            return
        p = para(doc)
        add_rich(p, ' · '.join(chips), grey(9, False))
        p.paragraph_format.space_after = Pt(8)
        chips.clear()

    for it in items:
        role, raw = it['role'], it['text']

        # Vain varsinaiset kysymykset liitteeseen — tunnistus tekstistä, ei
        # luokkanimestä (esim. oc-note on sisältöä, ei muistiinpano).
        if role == 'note' and NOTE_Q.search(raw):
            questions.append((where, raw))
            continue

        text = clean(raw, questions, where)
        if not text:
            continue

        if role == 'chip':
            chips.append(text)
            continue
        flush_chips()

        if role == 'step':
            add_rich(para(doc, 'Heading 4'), text)
        elif role == 'stat':
            add_rich(para(doc), text, bold_fmt)
        elif role == 'image':
            p = para(doc)
            add_rich(p, 'Kuva/grafiikka: ' + text, grey(9))
            p.paragraph_format.space_after = Pt(8)
        elif role == 'h3':
            add_rich(para(doc, 'Heading 3'), text)
        elif role == 'h4':
            add_rich(para(doc, 'Heading 4'), text)
        elif role == 'summary':
            add_rich(para(doc, 'Heading 4'), text)
        elif role == 'eyebrow':
            p = para(doc)
            add_rich(p, text.upper(), eyebrow_fmt)
            p.paragraph_format.space_after = Pt(0)
        elif role == 'lede':
            add_rich(para(doc), text, lede_fmt)
        elif role == 'li':
            add_rich(para(doc, 'List Bullet'), text)
        elif role == 'cta':
            p = para(doc)
            r = p.add_run('▸  ')
            r.font.color.rgb = BLUE
            r.font.bold = True
            add_rich(p, text, bold_fmt)
            r2 = p.add_run('   (nappi)')
            r2.font.size = Pt(9)
            r2.font.color.rgb = GREY
            r2.font.italic = True
        elif role == 'label':
            p = para(doc)
            r = p.add_run('Lomakekenttä: ')
            r.font.size = Pt(10)
            r.font.color.rgb = GREY
            r.font.italic = True
            add_rich(p, text, small_fmt)
            p.paragraph_format.space_after = Pt(2)
        elif role == 'select':
            p = para(doc)
            add_rich(p, 'Valinnat: ' + text, grey(10))
            p.paragraph_format.space_after = Pt(2)
        elif role == 'dt':
            add_rich(para(doc), text, bold_fmt)
        elif role == 'dd':
            p = add_rich(para(doc), text)
            p.paragraph_format.left_indent = Cm(0.6)
        elif role == 'quote':
            add_rich(para(doc, 'Intense Quote'), text)
        elif role == 'note':
            p = para(doc)
            r = p.add_run('Huom: ')
            r.font.size = Pt(9)
            r.font.bold = True
            r.font.italic = True
            r.font.color.rgb = GREY
            add_rich(p, text, grey(9))
        elif role == 'text' and len(text) < 50:
            p = para(doc)
            add_rich(p, text, grey(10, False))
            p.paragraph_format.space_after = Pt(2)
        else:
            add_rich(para(doc), text)

    flush_chips()


# ---------------------------------------------------------------- dokumentti
def build(tree, idx, outdir, stamp):
    doc = Document()
    setup_styles(doc)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.2)

    label = tree['label']
    questions = []

    # --- kansiosa
    p = para(doc)
    r = p.add_run('CONDITE.FI — SIVUSTOUUDISTUS')
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = BLUE
    p.paragraph_format.space_after = Pt(2)

    add_rich(para(doc, 'Title'), label)

    p = para(doc)
    add_rich(p, tree['title'], grey(10, False))
    p.paragraph_format.space_after = Pt(2)

    meta = [('Tiedosto', tree['file'])]
    if tree.get('meta_desc'):
        meta.append(('Meta-kuvaus', tree['meta_desc']))
    meta.append(('Poimittu', stamp))
    for k, v in meta:
        p = para(doc)
        r = p.add_run(k + ': ')
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = GREY
        r2 = p.add_run(v)
        r2.font.size = Pt(9)
        r2.font.color.rgb = GREY
        p.paragraph_format.space_after = Pt(1)

    p = para(doc)
    r = p.add_run('Keltaisella korostettu teksti on wireframe-luonnos '
                  'tai paikanpitäjä — ei vielä hyväksyttyä copya.')
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = MID
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)

    # --- sivun tekstit
    for s in tree['sections']:
        head = s['heading']
        items = list(s['items'])
        # sivulla eyebrow on otsikon YLÄPUOLELLA — renderöi se ensin
        lead = []
        while items and items[0]['role'] == 'eyebrow':
            lead.append(items.pop(0))
        if head:
            head_clean = clean(head, questions, head)
            where = head_clean
        else:
            head_clean, where = None, label
        if lead:
            render_items(doc, lead, questions, where)
        if head_clean:
            style = 'Heading 1' if s['level'] == 1 else 'Heading 2'
            add_rich(para(doc, style), head_clean)
        render_items(doc, items, questions, where)

    # --- sivupohjan elementit (navi, murupolku, footer)
    if tree.get('chrome'):
        doc.add_page_break()
        add_rich(para(doc, 'Heading 1'), 'Sivupohjan elementit')
        p = para(doc)
        r = p.add_run('Nämä toistuvat samanlaisina kaikilla sivuilla.')
        r.font.size = Pt(9)
        r.font.italic = True
        r.font.color.rgb = MID
        for c in tree['chrome']:
            add_rich(para(doc, 'Heading 2'), c['name'])
            render_items(doc, c['items'], questions, c['name'])

    # --- liite: avoimet kysymykset
    if questions:
        doc.add_page_break()
        add_rich(para(doc, 'Heading 1'), 'Liite: avoimet kysymykset')
        p = para(doc)
        r = p.add_run('Sisäisiä muistiinpanoja wireframesta — vahvistettavia tietoja '
                      'ja ratkaisemattomia sisältökysymyksiä. Eivät ole sivun tekstiä.')
        r.font.size = Pt(9)
        r.font.italic = True
        r.font.color.rgb = MID
        seen = set()
        for where, q in questions:
            key = q[:80]
            if key in seen:
                continue
            seen.add(key)
            p = para(doc, 'List Bullet')
            if where and where != q:
                r = p.add_run(where + ' — ')
                r.font.bold = True
                r.font.size = Pt(10)
            r2 = p.add_run(re.sub(r'\s+', ' ', q).strip())
            r2.font.size = Pt(10)

    name = f'{idx:02d}_{slug(label)}.docx'
    path = os.path.join(outdir, name)
    doc.save(path)
    return name, len(questions)


if __name__ == '__main__':
    trees = json.load(io.open(os.path.join(HERE, 'trees.json'), encoding='utf-8'))
    os.makedirs(OUT, exist_ok=True)
    stamp = datetime.date.today().strftime('%-d.%-m.%Y')

    # render.py:n järjestys
    sys.path.insert(0, HERE)
    from render import PAGES
    order = [f for f, _ in PAGES if f in trees]

    for i, f in enumerate(order, 1):
        name, nq = build(trees[f], i, OUT, stamp)
        size = os.path.getsize(os.path.join(OUT, name))
        print(f'  {name:44} {size/1024:6.1f} kt  {nq:>2} kysymystä liitteessä')
    print(f'\n→ {OUT} ({len(order)} dokumenttia)')
